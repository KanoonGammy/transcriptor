import io
from datetime import datetime
from typing import Dict, Any, List, Optional
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from app.elevenlabs_service import format_seconds

def set_cell_background(cell, color_hex: str):
    """Set background color for a table cell."""
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner padding for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def export_to_docx(transcription_data: Dict[str, Any]) -> io.BytesIO:
    """
    Generate a beautifully styled Microsoft Word document (.docx).
    """
    doc = Document()

    # Page setup: Normal 1-inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    title_text = transcription_data.get("title") or transcription_data.get("filename") or "เอกสารการถอดเสียง"
    filename = transcription_data.get("filename", "-")
    duration_sec = transcription_data.get("duration") or 0.0
    lang = transcription_data.get("language_code") or "th"
    model_id = transcription_data.get("model_id") or "scribe_v1"
    created_at = transcription_data.get("created_at", "")
    speaker_map = transcription_data.get("speaker_map") or {}
    segments: List[Dict[str, Any]] = transcription_data.get("segments") or []
    full_text = transcription_data.get("formatted_text") or transcription_data.get("full_text") or ""

    # Document Header / Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run(title_text)
    run_title.font.name = "Sarabun"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 41, 59) # Slate 800

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(16)
    run_sub = sub_p.add_run("เอกสารถอดเสียงอัตโนมัติ (ElevenLabs Scribe)")
    run_sub.font.name = "Sarabun"
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = RGBColor(100, 116, 139) # Slate 500

    # Metadata Info Table
    meta_table = doc.add_table(rows=2, cols=3)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_items = [
        ("📁 ไฟล์ต้นฉบับ", filename[:32]),
        ("⏱️ ความยาวเสียง", format_seconds(duration_sec)),
        ("🌐 ภาษาที่ตรวจพบ", lang.upper()),
        ("🤖 โมเดลที่ใช้", model_id),
        ("👥 จำนวนผู้พูด", f"{len(speaker_map)} คน" if speaker_map else "ไม่ระบุ"),
        ("📅 วันที่ถอดเสียง", created_at[:10] if created_at else datetime.now().strftime("%Y-%m-%d")),
    ]

    for idx, (label, val) in enumerate(meta_items):
        r = idx // 3
        c = idx % 3
        cell = meta_table.cell(r, c)
        set_cell_background(cell, "F8FAFC") # Slate 50
        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        
        lbl_run = p.add_run(f"{label}\n")
        lbl_run.font.name = "Sarabun"
        lbl_run.font.size = Pt(9)
        lbl_run.font.color.rgb = RGBColor(100, 116, 139)
        
        val_run = p.add_run(val)
        val_run.font.name = "Sarabun"
        val_run.font.size = Pt(10)
        val_run.font.bold = True
        val_run.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Heading for Dialogue Transcript
    sec_head = doc.add_paragraph()
    sec_head.paragraph_format.space_before = Pt(12)
    sec_head.paragraph_format.space_after = Pt(8)
    run_sec = sec_head.add_run("บทสนทนาที่ถอดความได้ (Transcript)")
    run_sec.font.name = "Sarabun"
    run_sec.font.size = Pt(15)
    run_sec.font.bold = True
    run_sec.font.color.rgb = RGBColor(15, 23, 42)

    # Palette for different speakers
    speaker_colors = [
        RGBColor(37, 99, 235),   # Blue 600
        RGBColor(13, 148, 136),  # Teal 600
        RGBColor(217, 119, 6),   # Amber 600
        RGBColor(147, 51, 234),  # Purple 600
        RGBColor(225, 29, 72),   # Rose 600
        RGBColor(79, 70, 229),   # Indigo 600
    ]
    speaker_color_map = {}
    color_idx = 0

    if segments:
        for seg in segments:
            spk_id = seg.get("speaker_id") or "speaker_0"
            spk_name = speaker_map.get(spk_id, spk_id)
            start_str = format_seconds(seg.get("start", 0.0))
            end_str = format_seconds(seg.get("end", 0.0))
            text_content = seg.get("text", "").strip()

            if spk_id not in speaker_color_map:
                speaker_color_map[spk_id] = speaker_colors[color_idx % len(speaker_colors)]
                color_idx += 1

            p_spk = doc.add_paragraph()
            p_spk.paragraph_format.space_before = Pt(10)
            p_spk.paragraph_format.space_after = Pt(2)
            p_spk.paragraph_format.line_spacing = 1.15

            # Timestamp badge
            time_run = p_spk.add_run(f"[{start_str} - {end_str}] ")
            time_run.font.name = "Sarabun"
            time_run.font.size = Pt(9.5)
            time_run.font.color.rgb = RGBColor(148, 163, 184) # Slate 400

            # Speaker Name
            spk_run = p_spk.add_run(f"{spk_name}:")
            spk_run.font.name = "Sarabun"
            spk_run.font.size = Pt(11)
            spk_run.font.bold = True
            spk_run.font.color.rgb = speaker_color_map[spk_id]

            # Dialogue body
            p_body = doc.add_paragraph()
            p_body.paragraph_format.left_indent = Inches(0.25)
            p_body.paragraph_format.space_before = Pt(0)
            p_body.paragraph_format.space_after = Pt(6)
            p_body.paragraph_format.line_spacing = 1.25

            body_run = p_body.add_run(text_content)
            body_run.font.name = "Sarabun"
            body_run.font.size = Pt(11)
            body_run.font.color.rgb = RGBColor(51, 65, 85) # Slate 700

    else:
        # Fallback to plain paragraphs
        paragraphs = full_text.split("\n")
        for para in paragraphs:
            if para.strip():
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.25
                r = p.add_run(para.strip())
                r.font.name = "Sarabun"
                r.font.size = Pt(11)
                r.font.color.rgb = RGBColor(51, 65, 85)

    # Output to BytesIO
    target_stream = io.BytesIO()
    doc.save(target_stream)
    target_stream.seek(0)
    return target_stream
